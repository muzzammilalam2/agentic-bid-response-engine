from fastapi import FastAPI, HTTPException, UploadFile, File, Form, Header, Depends, Query
from fastapi.responses import FileResponse
from pydantic import BaseModel
import uuid
import os
from dotenv import load_dotenv, find_dotenv
load_dotenv(find_dotenv())

import io
import docx
from docx.shared import Inches, Pt, RGBColor
from database import db_connect, init_db
from agents import ProposalWorkflowManager
from document_parser import parse_document, chunk_text
import json
import re

# Initialize SQLite database schema
init_db()

app = FastAPI(title="AI-Powered Bid & Proposal Response Engine API")
workflow_manager = ProposalWorkflowManager()

def check_permission(username: str, required_permission: str):
    conn = db_connect()
    cursor = conn.cursor()
    cursor.execute("SELECT permissions FROM users WHERE username = ?", (username,))
    row = cursor.fetchone()
    conn.close()
    if not row:
        # Fallback to viewer permissions if user does not exist
        conn = db_connect()
        cursor = conn.cursor()
        cursor.execute("SELECT permissions FROM users WHERE username = 'viewer'")
        row = cursor.fetchone()
        conn.close()
        username = "viewer"
        
    permissions = json.loads(row["permissions"])
    if required_permission not in permissions:
        raise HTTPException(
            status_code=403, 
            detail=f"Permission Denied: User '{username}' does not have the '{required_permission}' permission."
        )

def verify_permission(required_permission: str):
    def dependency(x_username: str = Header(None), username: str = Query(None)):
        active_user = x_username or username or "viewer"
        check_permission(active_user, required_permission)
        return active_user
    return dependency

class ChecklistOverride(BaseModel):
    status: str
    notes: str = None

class DraftEdit(BaseModel):
    drafted_answer: str

@app.get("/")
def read_root():
    return {"status": "healthy", "service": "Bid Proposal Response Engine API"}

@app.get("/workspaces")
def get_workspaces():
    conn = db_connect()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM workspaces")
    ws_rows = cursor.fetchall()
    
    results = []
    for ws_row in ws_rows:
        ws = dict(ws_row)
        ws_id = ws["id"]
        
        # Fetch checklist to calculate compliance percentage and gaps
        cursor.execute("SELECT status FROM compliance_checklist WHERE workspace_id = ?", (ws_id,))
        checklist_rows = cursor.fetchall()
        
        total_items = len(checklist_rows)
        passed_items = sum(1 for item in checklist_rows if item["status"] == "Pass")
        gaps_count = sum(1 for item in checklist_rows if item["status"] == "Fail")
        compliance_pct = (passed_items / total_items * 100) if total_items > 0 else 100
        
        competitors_str = ws.get("competitors") or ""
        competitor_list = [c.strip() for c in competitors_str.split(",") if c.strip()]
        competitor_count = len(competitor_list) if competitor_list else 0
        
        scoring_res = workflow_manager.scorer.calculate_score(
            sector=ws.get("client", "IT Services"),
            budget_str=ws.get("budget", "PKR 0M"),
            compliance_pct=compliance_pct,
            gaps_count=gaps_count,
            competitor_count=competitor_count
        )
        
        ws["win_probability"] = scoring_res["win_probability"]
        ws["decision"] = scoring_res["decision"]
        ws["compliance_pct"] = round(compliance_pct, 1)
        ws["gaps_count"] = gaps_count
        
        results.append(ws)
        
    conn.close()
    return results

@app.get("/criteria")
def get_criteria():
    conn = db_connect()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM evaluation_criteria")
    rows = cursor.fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.get("/criteria/{sector}")
def get_criteria_by_sector(sector: str):
    conn = db_connect()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM evaluation_criteria WHERE sector = ?", (sector,))
    rows = cursor.fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.get("/workspaces/{workspace_id}")
def get_workspace_detail(workspace_id: str):
    conn = db_connect()
    cursor = conn.cursor()
    
    # Workspace details
    cursor.execute("SELECT * FROM workspaces WHERE id = ?", (workspace_id,))
    ws_row = cursor.fetchone()
    if not ws_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Workspace not found")
        
    # Compliance checklist
    cursor.execute("SELECT * FROM compliance_checklist WHERE workspace_id = ?", (workspace_id,))
    checklist = [dict(r) for r in cursor.fetchall()]
    
    # Proposal drafts
    cursor.execute("SELECT * FROM proposal_drafts WHERE workspace_id = ?", (workspace_id,))
    drafts = [dict(r) for r in cursor.fetchall()]

    # SOW Plan details
    cursor.execute("SELECT * FROM sow_plans WHERE workspace_id = ?", (workspace_id,))
    sow_row = cursor.fetchone()
    
    conn.close()
    
    sow_plan = None
    if sow_row:
        try:
            sow_plan = {
                "milestones": json.loads(sow_row["milestones"]),
                "efforts": json.loads(sow_row["efforts"]),
                "costs": json.loads(sow_row["costs"])
            }
        except Exception as e:
            print("Error loading SOW Plan JSON:", e)
            
    # Run scoring dynamically to get win breakdown
    ws = dict(ws_row)
    gaps_count = sum(1 for item in checklist if item["status"] == "Fail")
    total_items = len(checklist)
    passed_items = sum(1 for item in checklist if item["status"] == "Pass")
    compliance_pct = (passed_items / total_items * 100) if total_items > 0 else 100
    
    competitors_str = ws.get("competitors") or ""
    competitor_list = [c.strip() for c in competitors_str.split(",") if c.strip()]
    competitor_count = len(competitor_list) if competitor_list else 0
    
    scoring_res = workflow_manager.scorer.calculate_score(
        sector=ws.get("client", "IT Services"), # fallbacks
        budget_str=ws.get("budget", "PKR 0M"),
        compliance_pct=compliance_pct,
        gaps_count=gaps_count,
        competitor_count=competitor_count
    )
    
    return {
        "workspace": ws,
        "compliance_checklist": checklist,
        "proposal_drafts": drafts,
        "sow_plan": sow_plan,
        "win_probability": scoring_res["win_probability"],
        "decision": scoring_res["decision"],
        "scoring_breakdown": scoring_res["breakdown"]
    }

@app.post("/workspaces")
def create_workspace(
    title: str = Form(...),
    sector: str = Form(...),
    budget: str = Form(...),
    deadline: str = Form(...),
    competitors: str = Form(None),
    rfp_text: str = Form(None),
    rfp_file: UploadFile = File(None),
    use_llm: bool = Form(False),
    llm_model: str = Form("gemini-1.5-flash"),
    x_gemini_api_key: str = Header(None),
    username: str = Depends(verify_permission("create_workspace"))
):
    workspace_id = str(uuid.uuid4())[:8]

    if rfp_file:
        try:
            content = rfp_file.file.read()
            # Save file to uploads folder on disk under workspace_id directory
            uploads_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "uploads", workspace_id)
            os.makedirs(uploads_dir, exist_ok=True)
            file_path = os.path.join(uploads_dir, rfp_file.filename)
            with open(file_path, "wb") as f:
                f.write(content)
                
            rfp_text = parse_document(content, rfp_file.filename)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse RFP file: {e}")
            
    if not rfp_text or not rfp_text.strip():
        raise HTTPException(status_code=400, detail="RFP requirements text or uploaded RFP file is required.")
    
    # Run Agentic Proposal Loop (handling both Offline and LLM modes)
    res = workflow_manager.run_workflow(
        workspace_id=workspace_id,
        rfp_title=title,
        rfp_text=rfp_text,
        sector=sector,
        budget_str=budget,
        deadline=deadline,
        use_llm=use_llm,
        api_key=x_gemini_api_key,
        llm_model=llm_model,
        competitors=competitors
    )
    
    return {
        "message": "Workspace initialized successfully",
        "workspace_id": workspace_id,
        "workflow_results": res
    }

@app.post("/capabilities/upload")
def upload_capability_document(
    file: UploadFile = File(...),
    domain: str = Form("General"),
    username: str = Depends(verify_permission("create_workspace"))
):
    try:
        content = file.file.read()
        # Save file to capabilities folder in root directory
        caps_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "capabilities")
        os.makedirs(caps_dir, exist_ok=True)
        file_path = os.path.join(caps_dir, file.filename)
        with open(file_path, "wb") as f:
            f.write(content)
            
        text = parse_document(content, file.filename)
        if not text.strip():
            raise HTTPException(status_code=400, detail="Uploaded file is empty or text extraction failed.")
            
        chunks = chunk_text(text, chunk_size=600, overlap=100)
        if not chunks:
            raise HTTPException(status_code=400, detail="No readable text chunks could be extracted.")
            
        conn = db_connect()
        cursor = conn.cursor()
        
        for idx, chunk in enumerate(chunks):
            chunk_id = f"cap-upload-{str(uuid.uuid4())[:8]}-{idx}"
            metadata = json.dumps({
                "cap_id": f"UPLOAD-{str(uuid.uuid4())[:4]}",
                "filename": file.filename
            })
            cursor.execute("""
            INSERT OR REPLACE INTO capability_chunks (chunk_id, source_file, content, domain, metadata)
            VALUES (?, ?, ?, ?, ?)
            """, (chunk_id, file.filename, chunk, domain, metadata))
            
        conn.commit()
        conn.close()
        
        # Hot-reload the RAG engine index to include the new chunks
        workflow_manager.rag.reload_index()
        
        return {
            "message": f"Ingested capability document '{file.filename}' successfully.",
            "chunks_count": len(chunks)
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/workspaces/{workspace_id}/checklist/{checklist_id}/override")
def override_checklist(
    workspace_id: str,
    checklist_id: str,
    payload: ChecklistOverride,
    username: str = Depends(verify_permission("override_checklist"))
):
    conn = db_connect()
    cursor = conn.cursor()
    
    cursor.execute("""
    UPDATE compliance_checklist 
    SET status = ?, notes = ?, override_status = 'Manual Override'
    WHERE id = ? AND workspace_id = ?
    """, (payload.status, payload.notes, checklist_id, workspace_id))
    
    conn.commit()
    conn.close()
    return {"message": "Compliance checklist status overridden successfully"}

@app.post("/workspaces/{workspace_id}/drafts/{question_id}/edit")
def edit_draft(
    workspace_id: str,
    question_id: str,
    payload: DraftEdit,
    username: str = Depends(verify_permission("edit_draft"))
):
    conn = db_connect()
    cursor = conn.cursor()
    
    cursor.execute("""
    UPDATE proposal_drafts 
    SET drafted_answer = ?
    WHERE question_id = ? AND workspace_id = ?
    """, (payload.drafted_answer, question_id, workspace_id))
    
    conn.commit()
    conn.close()
    return {"message": "Draft answer updated successfully"}

@app.get("/workspaces/{workspace_id}/export")
def export_workspace_proposal(
    workspace_id: str,
    username: str = Depends(verify_permission("export_proposal"))
):
    conn = db_connect()
    cursor = conn.cursor()
    
    # Workspace details
    cursor.execute("SELECT * FROM workspaces WHERE id = ?", (workspace_id,))
    ws = cursor.fetchone()
    if not ws:
        conn.close()
        raise HTTPException(status_code=404, detail="Workspace not found")
        
    # Proposal drafts
    cursor.execute("SELECT * FROM proposal_drafts WHERE workspace_id = ? ORDER BY question_id ASC", (workspace_id,))
    drafts = cursor.fetchall()
    
    conn.close()
    
    # Generate python-docx document
    doc = docx.Document()
    
    # Style styling
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Add Title Section
    title = doc.add_paragraph()
    run = title.add_run("BID & PROPOSAL TECHNICAL RESPONSE")
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    
    sub = doc.add_paragraph()
    r = sub.add_run(f"Project Workspace: {ws['title']}\nTarget Client Sector: {ws['client']}\nBudget Reference: {ws['budget']}\nSubmission Deadline: {ws['deadline']}")
    r.italic = True
    
    doc.add_paragraph().add_run("-" * 60)
    
    # Add Draft Sections
    for draft in drafts:
        p_title = doc.add_paragraph()
        p_title_run = p_title.add_run(f"Question {draft['question_id']}: {draft['question_text']}")
        p_title_run.bold = True
        p_title_run.font.size = Pt(14)
        p_title_run.font.color.rgb = RGBColor(43, 108, 176)
        
        # Clean markdown formatting from answer for Word
        clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', draft['drafted_answer'])
        clean_text = clean_text.replace("- ", "• ")
        
        doc.add_paragraph(clean_text)
        
        p_ref = doc.add_paragraph()
        pref_run = p_ref.add_run(f"Matched Evidence Context:\n{draft['matched_evidence']}")
        pref_run.font.size = Pt(9.5)
        pref_run.font.color.rgb = RGBColor(128, 128, 128)
        pref_run.italic = True
        
        doc.add_paragraph().add_run("-" * 40)
        
    # Fetch SOW Plan and append to Word document as Appendix
    conn = db_connect()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM sow_plans WHERE workspace_id = ?", (workspace_id,))
    sow_row = cursor.fetchone()
    conn.close()
    
    if sow_row:
        try:
            sow_milestones = json.loads(sow_row["milestones"])
            sow_efforts = json.loads(sow_row["efforts"])
            sow_costs = json.loads(sow_row["costs"])
            
            # Page Break & Appendix Title
            doc.add_page_break()
            sow_title = doc.add_paragraph()
            sow_title_run = sow_title.add_run("APPENDIX: STATEMENT OF WORK (SOW) PLAN")
            sow_title_run.bold = True
            sow_title_run.font.size = Pt(16)
            sow_title_run.font.color.rgb = RGBColor(31, 78, 121)
            
            # 1. Project Milestones Table
            doc.add_paragraph().add_run("1. Project Milestones & Deliverables Schedule").bold = True
            m_table = doc.add_table(rows=1, cols=3)
            m_table.style = 'Table Grid'
            hdr_cells = m_table.rows[0].cells
            hdr_cells[0].text = 'Phase'
            hdr_cells[1].text = 'Duration'
            hdr_cells[2].text = 'Deliverables'
            for cell in hdr_cells:
                cell.paragraphs[0].runs[0].font.bold = True
                
            for m in sow_milestones:
                row_cells = m_table.add_row().cells
                row_cells[0].text = str(m.get("phase", ""))
                row_cells[1].text = str(m.get("duration", ""))
                row_cells[2].text = str(m.get("deliverables", ""))
                
            doc.add_paragraph() # spacing
            
            # 2. Labor Effort Table
            doc.add_paragraph().add_run("2. Resource Headcount & Labor Effort Estimates").bold = True
            e_table = doc.add_table(rows=1, cols=5)
            e_table.style = 'Table Grid'
            hdr_cells = e_table.rows[0].cells
            hdr_cells[0].text = 'Role'
            hdr_cells[1].text = 'Count'
            hdr_cells[2].text = 'Man-Months'
            hdr_cells[3].text = 'Rate/Month'
            hdr_cells[4].text = 'Total Cost'
            for cell in hdr_cells:
                cell.paragraphs[0].runs[0].font.bold = True
                
            for e in sow_efforts:
                row_cells = e_table.add_row().cells
                row_cells[0].text = str(e.get("role", ""))
                row_cells[1].text = str(e.get("count", ""))
                row_cells[2].text = str(e.get("man_months", ""))
                row_cells[3].text = str(e.get("rate_per_month", ""))
                row_cells[4].text = str(e.get("total_cost", ""))
                
            doc.add_paragraph() # spacing
            
            # 3. Cost Analysis Table
            doc.add_paragraph().add_run("3. SOW Project Cost Analysis").bold = True
            c_table = doc.add_table(rows=1, cols=3)
            c_table.style = 'Table Grid'
            hdr_cells = c_table.rows[0].cells
            hdr_cells[0].text = 'Budget Category'
            hdr_cells[1].text = 'Estimated Cost (PKR)'
            hdr_cells[2].text = 'Allocation %'
            for cell in hdr_cells:
                cell.paragraphs[0].runs[0].font.bold = True
                
            for c in sow_costs:
                row_cells = c_table.add_row().cells
                row_cells[0].text = str(c.get("category", ""))
                row_cells[1].text = str(c.get("amount", ""))
                row_cells[2].text = str(c.get("percentage", ""))
                
        except Exception as ex:
            print("Error appending SOW Appendix to Word document:", ex)

    # Save document in memory
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    # Write to a temp file and return
    temp_dir = os.path.join(os.path.dirname(__file__), "temp")
    os.makedirs(temp_dir, exist_ok=True)
    temp_file_path = os.path.join(temp_dir, f"{workspace_id}_proposal.docx")
    
    with open(temp_file_path, "wb") as f:
        f.write(file_stream.read())
        
    return FileResponse(
        temp_file_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"{ws['title'].lower().replace(' ', '_')}_proposal.docx"
    )

class ChatRequest(BaseModel):
    question: str

@app.get("/workspaces/{workspace_id}/battlecard")
def get_battlecard(
    workspace_id: str,
    x_gemini_api_key: str = Header(None),
    use_llm: bool = Query(False),
    llm_model: str = Query("gemini-1.5-flash"),
    force_refresh: bool = Query(False)
):
    conn = db_connect()
    cursor = conn.cursor()
    
    if not force_refresh:
        cursor.execute("SELECT content FROM competitor_battlecards WHERE workspace_id = ?", (workspace_id,))
        row = cursor.fetchone()
        
        if row:
            conn.close()
            return {"battlecard": row["content"]}
        
    # Generate dynamically if missing
    cursor.execute("SELECT client, competitors FROM workspaces WHERE id = ?", (workspace_id,))
    ws_row = cursor.fetchone()
    if not ws_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Workspace not found")
        
    competitors = ws_row["competitors"]
    sector = ws_row["client"]
    conn.close()
    
    if not competitors or not competitors.strip():
        return {"battlecard": "No competitors specified for this workspace."}
        
    content = workflow_manager.generate_battlecard(
        workspace_id=workspace_id,
        sector=sector,
        competitors_str=competitors,
        use_llm=use_llm,
        api_key=x_gemini_api_key,
        llm_model=llm_model
    )
    return {"battlecard": content}

@app.post("/workspaces/{workspace_id}/chat")
def chat_workspace(
    workspace_id: str,
    payload: ChatRequest,
    x_gemini_api_key: str = Header(None),
    use_llm: bool = Query(False),
    llm_model: str = Query("gemini-1.5-flash")
):
    conn = db_connect()
    cursor = conn.cursor()
    cursor.execute("SELECT title, client, budget, deadline FROM workspaces WHERE id = ?", (workspace_id,))
    ws_row = cursor.fetchone()
    if not ws_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Workspace not found")
        
    cursor.execute("SELECT description FROM compliance_checklist WHERE workspace_id = ?", (workspace_id,))
    reqs = [r["description"] for r in cursor.fetchall()]
    
    cursor.execute("SELECT question_text, drafted_answer FROM proposal_drafts WHERE workspace_id = ?", (workspace_id,))
    drafts = [f"Q: {d['question_text']}\nDraft Answer: {d['drafted_answer']}" for d in cursor.fetchall()]
    conn.close()
    
    workspace_summary = (
        f"Workspace Title: {ws_row['title']}\n"
        f"Industry Sector: {ws_row['client']}\n"
        f"Budget: {ws_row['budget']}\n"
        f"Deadline: {ws_row['deadline']}\n"
        f"Extracted Requirements:\n" + "\n".join(f"- {r}" for r in reqs) + "\n\n"
        f"Proposal Draft Sections:\n" + "\n\n".join(drafts)
    )
    
    answer = ""
    if use_llm:
        effective_key = x_gemini_api_key or os.environ.get("GEMINI_API_KEY")
        if effective_key:
            try:
                import google.generativeai as genai
                genai.configure(api_key=effective_key)
                model = genai.GenerativeModel(llm_model)
                prompt = (
                    "You are a helpful bid-proposal workspace assistant/copilot.\n"
                    "You have access to the details of the active proposal workspace:\n"
                    f"=== WORKSPACE DETAILS ===\n{workspace_summary}\n=========================\n\n"
                    f"The user is asking: '{payload.question}'\n"
                    "Answer the user's question precisely and professionally based on the context above. If they ask about payment terms, timeline, or qualifications that are not explicitly documented, advise them clearly."
                )
                response = model.generate_content(prompt)
                answer = response.text.strip()
            except Exception as e:
                answer = f"Gemini Copilot encountered an error: {e}"
                
    if not answer:
        q = payload.question.lower()
        if "budget" in q or "price" in q:
            answer = f"The current workspace budget reference is {ws_row['budget']}."
        elif "deadline" in q or "due" in q:
            answer = f"The submission deadline is configured as {ws_row['deadline']}."
        elif "sector" in q or "vertical" in q:
            answer = f"This bid is configured under the '{ws_row['client']}' vertical."
        elif "requirements" in q or "compliance" in q:
            answer = f"The workspace has {len(reqs)} compliance checklist requirements:\n" + "\n".join(f"- {r}" for r in reqs[:5])
        else:
            answer = (
                "RFP Copilot (Offline Mode):\n"
                "To chat dynamically and query custom details, please enable **Gemini AI Mode** in the sidebar. "
                f"The current workspace context has {len(reqs)} requirements and {len(drafts)} proposal draft sections."
            )
            
    return {"answer": answer}

@app.post("/workspaces/{workspace_id}/checklist/{checklist_id}/auto-resolve")
def auto_resolve_checklist_item(
    workspace_id: str,
    checklist_id: str,
    x_gemini_api_key: str = Header(None),
    use_llm: bool = Query(False),
    llm_model: str = Query("gemini-1.5-flash")
):
    conn = db_connect()
    cursor = conn.cursor()
    cursor.execute("SELECT description, status FROM compliance_checklist WHERE id = ? AND workspace_id = ?", (checklist_id, workspace_id))
    row = cursor.fetchone()
    conn.close()
    
    if not row:
        raise HTTPException(status_code=404, detail="Checklist item not found")
        
    if row["status"] == "Pass":
        return {"message": "Requirement is already passing.", "status": "Pass"}
        
    res = workflow_manager.resolve_requirement_gap(
        workspace_id=workspace_id,
        checklist_id=checklist_id,
        requirement_desc=row["description"],
        use_llm=use_llm,
        api_key=x_gemini_api_key,
        llm_model=llm_model
    )
    return res
